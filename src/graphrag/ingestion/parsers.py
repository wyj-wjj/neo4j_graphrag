"""Format-specific parsers that preserve source locations and never execute active content."""

from __future__ import annotations

import asyncio
import tempfile
from pathlib import Path

import fitz
from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation

from graphrag.domain.errors import ValidationError
from graphrag.domain.models import Page, ParsedDocument
from graphrag.domain.ports import OCRPort


class ParserRegistry:
    def __init__(
        self, *, max_pages: int, ocr: OCRPort | None = None, ocr_timeout: float = 60
    ) -> None:
        self.max_pages = max_pages
        self.ocr = ocr
        self.ocr_timeout = ocr_timeout

    async def parse(self, filename: str, content: bytes) -> ParsedDocument:
        suffix = Path(filename).suffix.lower()
        if suffix in {".txt", ".md", ".html", ".htm"}:
            return self._parse_text(suffix, content)
        return await asyncio.to_thread(self._parse_binary, suffix, content)

    def _parse_text(self, suffix: str, content: bytes) -> ParsedDocument:
        text = content.decode("utf-8")
        if suffix in {".html", ".htm"}:
            soup = BeautifulSoup(text, "lxml")
            for node in soup(["script", "style", "iframe", "object", "embed"]):
                node.decompose()
            for tag in soup.find_all(True):
                for attribute in list(tag.attrs):
                    if attribute.lower().startswith("on"):
                        del tag.attrs[attribute]
            text = soup.get_text("\n", strip=True)
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        if not lines:
            raise ValidationError("文档没有可提取文本")
        pages = tuple(
            Page(
                text=line,
                page_number=1,
                title_path=(line.lstrip("# "),) if line.startswith("#") else (),
                block_type="heading" if line.startswith("#") else "paragraph",
            )
            for line in lines
        )
        return ParsedDocument(pages=pages, parser="html" if suffix in {".html", ".htm"} else "text")

    def _parse_binary(self, suffix: str, content: bytes) -> ParsedDocument:
        with tempfile.NamedTemporaryFile(suffix=suffix) as handle:
            handle.write(content)
            handle.flush()
            path = Path(handle.name)
            if suffix == ".pdf":
                return self._parse_pdf(path)
            if suffix == ".docx":
                return self._parse_docx(path)
            if suffix == ".xlsx":
                return self._parse_xlsx(path)
            if suffix == ".pptx":
                return self._parse_pptx(path)
            if suffix in {".png", ".jpg", ".jpeg"}:
                return self._parse_image(path)
        raise ValidationError("没有可用的文档解析器")

    def _parse_pdf(self, path: Path) -> ParsedDocument:
        try:
            document = fitz.open(path)
        except Exception as exc:
            raise ValidationError("PDF 文件损坏或无法打开") from exc
        if document.is_encrypted:
            document.close()
            raise ValidationError("不支持加密 PDF")
        if document.page_count > self.max_pages:
            document.close()
            raise ValidationError("PDF 页数超过限制")
        pages = []
        for number, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            if text:
                pages.append(Page(text=text, page_number=number))
        requires_ocr = not pages
        document.close()
        if requires_ocr:
            return ParsedDocument(
                pages=(Page(text="[等待 OCR]", page_number=1, block_type="image"),),
                parser="pymupdf",
                requires_ocr=True,
            )
        return ParsedDocument(pages=tuple(pages), parser="pymupdf")

    def _parse_docx(self, path: Path) -> ParsedDocument:
        try:
            document = Document(str(path))
        except Exception as exc:
            raise ValidationError("Word 文件损坏或无法打开") from exc
        pages: list[Page] = []
        title_path: tuple[str, ...] = ()
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            is_heading = bool(
                paragraph.style is not None and paragraph.style.name.startswith("Heading")
            )
            if is_heading:
                title_path = (text,)
            pages.append(
                Page(
                    text=text,
                    page_number=1,
                    title_path=title_path,
                    block_type="heading" if is_heading else "paragraph",
                )
            )
        for table in document.tables:
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
            if any(rows):
                pages.append(Page(text="\n".join(rows), page_number=1, block_type="table"))
        if not pages:
            raise ValidationError("Word 文档没有可提取内容")
        return ParsedDocument(pages=tuple(pages), parser="python-docx")

    def _parse_xlsx(self, path: Path) -> ParsedDocument:
        try:
            workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        except Exception as exc:
            raise ValidationError("Excel 文件损坏或无法打开") from exc
        pages: list[Page] = []
        total_cells = 0
        for sheet in workbook.worksheets:
            if sheet.sheet_state != "visible":
                continue
            rows = []
            for row in sheet.iter_rows(values_only=True):
                total_cells += len(row)
                if total_cells > 50_000:
                    workbook.close()
                    raise ValidationError("Excel 单元格数量超过限制")
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    rows.append(" | ".join(values))
            if rows:
                pages.append(
                    Page(
                        text="\n".join(rows),
                        page_number=len(pages) + 1,
                        title_path=(sheet.title,),
                        block_type="table",
                    )
                )
        workbook.close()
        if not pages:
            raise ValidationError("Excel 工作簿没有可见内容")
        return ParsedDocument(pages=tuple(pages), parser="openpyxl")

    def _parse_pptx(self, path: Path) -> ParsedDocument:
        try:
            presentation = Presentation(str(path))
        except Exception as exc:
            raise ValidationError("PowerPoint 文件损坏或无法打开") from exc
        if len(presentation.slides) > self.max_pages:
            raise ValidationError("PowerPoint 页数超过限制")
        pages: list[Page] = []
        for number, slide in enumerate(presentation.slides, start=1):
            text_parts: list[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    text = shape.text.strip()
                    if text:
                        text_parts.append(text)
                if getattr(shape, "has_table", False):
                    rows = [
                        " | ".join(cell.text.strip() for cell in row.cells)
                        for row in shape.table.rows
                    ]
                    text_parts.extend(rows)
            if text_parts:
                pages.append(Page(text="\n".join(text_parts), page_number=number))
        if not pages:
            raise ValidationError("PowerPoint 没有可提取内容")
        return ParsedDocument(pages=tuple(pages), parser="python-pptx")

    def _parse_image(self, path: Path) -> ParsedDocument:
        try:
            with Image.open(path) as image:
                image.verify()
        except Exception as exc:
            raise ValidationError("图片损坏或格式不匹配") from exc
        return ParsedDocument(
            pages=(Page(text="[等待 OCR]", page_number=1, block_type="image"),),
            parser="pillow",
            requires_ocr=True,
        )

    async def apply_ocr(self, parsed: ParsedDocument, content: bytes) -> ParsedDocument:
        if not parsed.requires_ocr:
            return parsed
        if self.ocr is None:
            raise ValidationError("文档需要 OCR，但未配置 OCR Adapter")
        pages = []
        for page in parsed.pages:
            text = await self.ocr.recognize(
                content, page_number=page.page_number, timeout=self.ocr_timeout
            )
            pages.append(page.model_copy(update={"text": text}))
        return ParsedDocument(pages=tuple(pages), parser=f"{parsed.parser}+ocr")
