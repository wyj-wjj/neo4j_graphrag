"""Deterministic multi-format knowledge fixtures and independent parsers."""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Callable
from datetime import datetime
from pathlib import Path
from typing import Literal

import fitz
from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from graphrag_data_factory.constants import DATASET_VERSION, SCHEMA_VERSION, SOURCE
from graphrag_data_factory.deterministic import deterministic_id, sha256_bytes
from graphrag_data_factory.models import KnowledgeRecord, PhysicalKnowledgeFileRecord

FileFormat = Literal["md", "txt", "html", "pdf", "docx", "xlsx", "pptx", "png", "jpeg"]
Variant = Literal["normal", "boundary", "corrupt"]

FORMATS: tuple[FileFormat, ...] = (
    "md",
    "txt",
    "html",
    "pdf",
    "docx",
    "xlsx",
    "pptx",
    "png",
    "jpeg",
)
VARIANTS: tuple[Variant, ...] = ("normal", "boundary", "corrupt")
_FIXED_OFFICE_TIMESTAMP = b"2025-01-01T00:00:00Z"
_MODIFIED_TIMESTAMP = re.compile(rb"(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)")
MIME_TYPES: dict[FileFormat, str] = {
    "md": "text/markdown",
    "txt": "text/plain",
    "html": "text/html",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "png": "image/png",
    "jpeg": "image/jpeg",
}


class PhysicalKnowledgeFileFactory:
    """Create nine supported formats with normal, boundary and corrupt variants."""

    def __init__(self, dataset_id: str) -> None:
        self.dataset_id = dataset_id

    def generate(
        self, knowledge: tuple[KnowledgeRecord, ...], *, count: int = 27
    ) -> tuple[tuple[PhysicalKnowledgeFileRecord, ...], dict[str, bytes]]:
        if not knowledge:
            raise ValueError("physical knowledge generation requires knowledge records")
        if count < len(FORMATS) * len(VARIANTS):
            raise ValueError("physical knowledge count must cover the complete format matrix")
        records: list[PhysicalKnowledgeFileRecord] = []
        files: dict[str, bytes] = {}
        for index in range(count):
            file_format = FORMATS[(index // len(VARIANTS)) % len(FORMATS)]
            variant = VARIANTS[index % len(VARIANTS)]
            document = knowledge[index % len(knowledge)]
            anchor = document.evidence_anchor_id
            content = self._render(file_format, variant, anchor)
            stem = variant if index < 27 else f"{variant}-{index:06d}"
            relative_path = f"knowledge-formats/{file_format}/{stem}.{file_format}"
            files[relative_path] = content
            records.append(
                PhysicalKnowledgeFileRecord(
                    dataset_id=self.dataset_id,
                    dataset_version=DATASET_VERSION,
                    scenario_id=f"scenario-physical-knowledge-{index:06d}",
                    tenant_id=document.tenant_id,
                    source=SOURCE,
                    is_synthetic=True,
                    schema_version=SCHEMA_VERSION,
                    physical_file_id=deterministic_id(
                        self.dataset_id, "physical-knowledge", str(index)
                    ),
                    document_id=document.document_id,
                    evidence_anchor_id=None if variant == "corrupt" else anchor,
                    format=file_format,
                    variant=variant,
                    relative_path=relative_path,
                    mime_type=MIME_TYPES[file_format],
                    size_bytes=len(content),
                    content_sha256=sha256_bytes(content),
                    expected_outcome=(
                        "reject_corrupt"
                        if variant == "corrupt"
                        else "ocr_success"
                        if file_format in {"png", "jpeg"}
                        else "parse_success"
                    ),
                )
            )
        return tuple(records), files

    @staticmethod
    def _render(file_format: FileFormat, variant: Variant, anchor: str) -> bytes:
        if variant == "corrupt":
            return b"\xff\xfeCORRUPT-SYNTHETIC-" + file_format.encode("ascii")
        boundary = variant == "boundary"
        text = _fixture_text(anchor, boundary)
        renderers: dict[FileFormat, Callable[[str, bool], bytes]] = {
            "md": _render_markdown,
            "txt": _render_text,
            "html": _render_html,
            "pdf": _render_pdf,
            "docx": _render_docx,
            "xlsx": _render_xlsx,
            "pptx": _render_pptx,
            "png": lambda value, _: _render_image(value, "PNG"),
            "jpeg": lambda value, _: _render_image(value, "JPEG"),
        }
        return renderers[file_format](text, boundary)


def parse_physical_file(path: Path, file_format: FileFormat) -> str:
    """Parse a physical fixture without importing any production ingestion code."""

    if file_format in {"md", "txt", "html"}:
        return path.read_text(encoding="utf-8")
    if file_format == "pdf":
        with fitz.open(path) as document:
            return "\n".join(page.get_text() for page in document)
    if file_format == "docx":
        return "\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs)
    if file_format == "xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return "\n".join(
                str(cell)
                for worksheet in workbook.worksheets
                for row in worksheet.iter_rows(values_only=True)
                for cell in row
                if cell is not None
            )
        finally:
            workbook.close()
    if file_format == "pptx":
        presentation = Presentation(str(path))
        return "\n".join(
            shape.text
            for slide in presentation.slides
            for shape in slide.shapes
            if hasattr(shape, "text")
        )
    with Image.open(path) as image:
        image.verify()
    return "verified-image"


def is_physical_file_rejected(path: Path, file_format: FileFormat) -> bool:
    """Return whether an independent parser rejects an intentionally corrupt fixture."""

    try:
        parse_physical_file(path, file_format)
    except Exception:
        return True
    return False


def _fixture_text(anchor: str, boundary: bool) -> str:
    suffix = (
        "\nBOUNDARY: unicode=中文/emoji-not-required; table=100x1; long=" + "X" * 4096
        if boundary
        else ""
    )
    return (
        "Synthetic/Fake knowledge fixture. Never use for real business.\n"
        f"evidence_anchor:{anchor}\n"
        "Rule SYN-PHYSICAL-001 applies only to this deterministic test dataset."
        f"{suffix}"
    )


def _render_markdown(text: str, boundary: bool) -> bytes:
    heading = "# Synthetic boundary fixture" if boundary else "# Synthetic fixture"
    return f"{heading}\n\n{text}\n".encode()


def _render_text(text: str, _: bool) -> bytes:
    return (text + "\n").encode()


def _render_html(text: str, boundary: bool) -> bytes:
    title = "Boundary" if boundary else "Normal"
    escaped = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return (
        '<!doctype html><html><head><meta charset="utf-8"><title>'
        f"Synthetic {title}</title></head><body><pre>{escaped}</pre></body></html>\n"
    ).encode()


def _render_pdf(text: str, _: bool) -> bytes:
    printable = text.encode("ascii", "replace").decode("ascii").replace("\\", "\\\\")
    printable = printable.replace("(", "\\(").replace(")", "\\)").replace("\n", " ")
    stream = f"BT /F1 8 Tf 36 760 Td ({printable}) Tj ET".encode("ascii")
    objects = (
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length "
        + str(len(stream)).encode("ascii")
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    )
    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for index, payload in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode("ascii"))
        output.extend(payload)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    return bytes(output)


def _render_docx(text: str, boundary: bool) -> bytes:
    document = Document()
    _set_core_properties(document.core_properties)
    document.add_heading("Synthetic boundary fixture" if boundary else "Synthetic fixture", 0)
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return _normalize_zip(buffer.getvalue())


def _render_xlsx(text: str, boundary: bool) -> bytes:
    workbook = Workbook()
    _set_core_properties(workbook.properties)
    worksheet = workbook.active
    worksheet.title = "Synthetic"
    worksheet["A1"] = "Synthetic boundary fixture" if boundary else "Synthetic fixture"
    worksheet["A2"] = text
    buffer = io.BytesIO()
    workbook.save(buffer)
    return _normalize_zip(buffer.getvalue())


def _render_pptx(text: str, boundary: bool) -> bytes:
    presentation = Presentation()
    _set_core_properties(presentation.core_properties)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(6))
    box.text = ("Synthetic boundary fixture\n" if boundary else "Synthetic fixture\n") + text
    buffer = io.BytesIO()
    presentation.save(buffer)
    return _normalize_zip(buffer.getvalue())


def _render_image(text: str, image_format: Literal["PNG", "JPEG"]) -> bytes:
    image = Image.new("RGB", (1600, 240), color="white")
    ImageDraw.Draw(image).text((20, 20), text[:600], fill="black")
    buffer = io.BytesIO()
    if image_format == "PNG":
        image.save(buffer, format=image_format, compress_level=9)
    else:
        image.save(buffer, format=image_format, quality=90, subsampling=0, optimize=False)
    return buffer.getvalue()


def _set_core_properties(properties: object) -> None:
    fixed = datetime(2025, 1, 1, 0, 0, 0)
    for name, value in (
        ("title", "Synthetic/Fake fixture"),
        ("creator", "graphrag-data-factory"),
        ("last_modified_by", "graphrag-data-factory"),
        ("created", fixed),
        ("modified", fixed),
    ):
        if hasattr(properties, name):
            setattr(properties, name, value)


def _normalize_zip(content: bytes) -> bytes:
    source = io.BytesIO(content)
    target = io.BytesIO()
    with (
        zipfile.ZipFile(source, "r") as archive,
        zipfile.ZipFile(
            target, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
        ) as normalized,
    ):
        for name in sorted(archive.namelist()):
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            payload = archive.read(name)
            if name == "docProps/core.xml":
                payload = _MODIFIED_TIMESTAMP.sub(
                    rb"\g<1>" + _FIXED_OFFICE_TIMESTAMP + rb"\g<2>", payload
                )
            normalized.writestr(info, payload)
    return target.getvalue()


__all__ = [
    "FORMATS",
    "MIME_TYPES",
    "PhysicalKnowledgeFileFactory",
    "is_physical_file_rejected",
    "parse_physical_file",
]
