from __future__ import annotations

import io
import zipfile
from pathlib import Path

import fitz
import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation

from graphrag.domain.errors import ValidationError
from graphrag.domain.models import Page, ParsedDocument
from graphrag.infrastructure.fakes import FakeOCR
from graphrag.ingestion.chunker import StructureAwareChunker
from graphrag.ingestion.parsers import ParserRegistry
from graphrag.ingestion.validation import UploadValidator


def validator(*, max_bytes: int = 1024) -> UploadValidator:
    return UploadValidator(
        allowed_extensions=frozenset({".txt", ".html", ".pdf", ".docx"}),
        max_bytes=max_bytes,
    )


def test_upload_validation_accepts_text_and_rejects_traversal_signature_and_size() -> None:
    result = validator().validate("policy.txt", "保修政策".encode())
    assert result.mime_type == "text/plain"
    with pytest.raises(ValidationError):
        validator().validate("../policy.txt", b"data")
    with pytest.raises(ValidationError):
        validator().validate("fake.pdf", b"not-pdf")
    with pytest.raises(ValidationError):
        validator(max_bytes=2).validate("large.txt", b"123")


def test_office_zip_bomb_and_path_traversal_are_rejected() -> None:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("../evil.xml", "x")
    with pytest.raises(ValidationError):
        validator(max_bytes=4096).validate("evil.docx", output.getvalue())


@pytest.mark.asyncio
async def test_text_html_and_ocr_parsers() -> None:
    parser = ParserRegistry(max_pages=10, ocr=FakeOCR())
    text = await parser.parse("notes.md", b"# Title\nBody")
    assert text.pages[0].block_type == "heading"
    html = await parser.parse(
        "page.html", b"<h1>Safe</h1><script>secret()</script><p onclick='x()'>Body</p>"
    )
    combined = " ".join(item.text for item in html.pages)
    assert "secret" not in combined
    image = b"\x89PNG\r\n\x1a\n" + b"not-a-real-image"
    parsed = ParsedDocument(
        pages=(Page(text="[waiting]", page_number=1, block_type="image"),),
        parser="test",
        requires_ocr=True,
    )
    recognized = await parser.apply_ocr(parsed, image)
    assert "Fake OCR" in recognized.pages[0].text


def test_structure_aware_chunking_preserves_tables_and_parent_context() -> None:
    parsed = ParsedDocument(
        pages=(
            Page(text="Policy", page_number=1, title_path=("Policy",), block_type="heading"),
            Page(text="A" * 260, page_number=1),
            Page(text="a | b\n1 | 2", page_number=2, block_type="table"),
        ),
        parser="test",
    )
    chunks = StructureAwareChunker(target_chars=100, max_chars=120, overlap_ratio=0.1).split(parsed)
    assert len(chunks) >= 4
    assert any(item.parent_content is not None for item in chunks)
    assert any("table" in item.source_location for item in chunks)
    assert [item.ordinal for item in chunks] == list(range(len(chunks)))


@pytest.mark.asyncio
async def test_all_binary_parsers_extract_synthetic_documents(tmp_path: Path) -> None:
    parser = ParserRegistry(max_pages=10, ocr=FakeOCR())

    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Warranty policy")
    pdf_bytes = pdf.tobytes()
    pdf.close()
    assert "Warranty" in (await parser.parse("policy.pdf", pdf_bytes)).pages[0].text

    word = Document()
    word.add_heading("Policy", level=1)
    word.add_paragraph("Two year warranty")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Term"
    table.cell(0, 1).text = "Two years"
    word_buffer = io.BytesIO()
    word.save(word_buffer)
    parsed_word = await parser.parse("policy.docx", word_buffer.getvalue())
    assert any(item.block_type == "table" for item in parsed_word.pages)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rules"
    sheet.append(["Item", "Value"])
    sheet.append(["Warranty", 2])
    excel_buffer = io.BytesIO()
    workbook.save(excel_buffer)
    workbook.close()
    parsed_excel = await parser.parse("policy.xlsx", excel_buffer.getvalue())
    assert parsed_excel.pages[0].title_path == ("Rules",)

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Policy"
    slide.placeholders[1].text = "Two year warranty"
    ppt_buffer = io.BytesIO()
    presentation.save(ppt_buffer)
    parsed_ppt = await parser.parse("policy.pptx", ppt_buffer.getvalue())
    assert "Two year" in parsed_ppt.pages[0].text

    image_buffer = io.BytesIO()
    Image.new("RGB", (4, 4), "white").save(image_buffer, format="PNG")
    image = await parser.parse("scan.png", image_buffer.getvalue())
    assert image.requires_ocr
    assert "Fake OCR" in (await parser.apply_ocr(image, image_buffer.getvalue())).pages[0].text


@pytest.mark.asyncio
async def test_parser_rejects_corrupt_and_empty_documents() -> None:
    parser = ParserRegistry(max_pages=1)
    with pytest.raises(ValidationError):
        await parser.parse("empty.txt", b"\n")
    with pytest.raises(ValidationError):
        await parser.parse("bad.pdf", b"%PDF-broken")
    with pytest.raises(ValidationError):
        await parser.parse("unknown.bin", b"data")
